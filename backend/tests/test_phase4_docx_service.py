"""Unit tests for Phase 4: DOCX Service (Parsing and Exporting)."""

from pathlib import Path
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
import pytest

from app.models.document import Document, DocumentMetadata, DocumentStyle, Page
from app.models.elements import (
    BoundingBox,
    TableCell,
    TableElement,
    TextAlignment,
    TextElement,
)
from app.services.docx.docx_service import DOCXService


@pytest.fixture
def sample_docx_file(tmp_path) -> Path:
    """Creates a sample DOCX with formatted paragraphs, alignments, and a table."""
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()

    # Paragraph 1: Centered title
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    r1.bold = True
    r1.font.size = Pt(14)

    # Paragraph 2: Normal body
    p2 = doc.add_paragraph("Đây là văn bản hợp đồng chính thức giữa hai bên.")

    # Table: 2x2
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "STT"
    table.cell(0, 1).text = "Hạng mục"
    table.cell(1, 0).text = "01"
    table.cell(1, 1).text = "Phát triển phần mềm Document Assistant"

    doc.save(str(docx_path))
    return docx_path


def test_docx_parsing(sample_docx_file):
    """Validates parsing a DOCX file into the intermediate Document model."""
    service = DOCXService()
    doc_model = service.parse_to_document(sample_docx_file)

    assert doc_model.total_pages == 1
    assert doc_model.metadata.source_format == "docx"

    page = doc_model.pages[0]
    text_elements = page.get_text_elements()
    table_elements = page.get_tables()

    # Check paragraphs
    assert any("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in el.text and el.bold for el in text_elements)
    assert any("Đây là văn bản hợp đồng" in el.text for el in text_elements)

    # Check table
    assert len(table_elements) == 1
    tab = table_elements[0]
    assert tab.rows == 2
    assert tab.columns == 2
    matrix = tab.to_matrix()
    assert matrix[0] == ["STT", "Hạng mục"]
    assert matrix[1] == ["01", "Phát triển phần mềm Document Assistant"]


def test_document_model_to_docx_export(tmp_path):
    """Validates exporting a Document model into a formatted DOCX."""
    doc_model = Document(
        metadata=DocumentMetadata(title="Test Export Doc"),
        styles=DocumentStyle(default_font="Arial", default_font_size=12.0),
    )

    page1 = Page(page_number=1, width=595.0, height=842.0)
    page1.add_element(
        TextElement(
            bbox=BoundingBox(x1=50, y1=50, x2=400, y2=80),
            text="BÁO CÁO TÀI CHÍNH QUÝ 1",
            font_size=16.0,
            bold=True,
            alignment=TextAlignment.CENTER,
        )
    )

    cells = [
        TableCell(row_index=0, col_index=0, text="Doanh thu"),
        TableCell(row_index=0, col_index=1, text="100 tỷ"),
        TableCell(row_index=1, col_index=0, text="Lợi nhuận"),
        TableCell(row_index=1, col_index=1, text="25 tỷ"),
    ]
    page1.add_element(
        TableElement(
            bbox=BoundingBox(x1=50, y1=100, x2=500, y2=200),
            rows=2,
            columns=2,
            cells=cells,
            has_header=True,
        )
    )

    doc_model.add_page(page1)

    out_docx = tmp_path / "exported.docx"
    service = DOCXService()
    service.export_to_docx(doc_model, out_docx)

    assert out_docx.exists()
    assert out_docx.stat().st_size > 0

    # Read back with python-docx to verify contents
    exported_doc = docx.Document(str(out_docx))
    all_paras = [p.text for p in exported_doc.paragraphs]
    assert any("BÁO CÁO TÀI CHÍNH QUÝ 1" in t for t in all_paras)

    assert len(exported_doc.tables) == 1
    exp_table = exported_doc.tables[0]
    assert exp_table.cell(0, 0).text == "Doanh thu"
    assert exp_table.cell(0, 1).text == "100 tỷ"
    assert exp_table.cell(1, 0).text == "Lợi nhuận"
    assert exp_table.cell(1, 1).text == "25 tỷ"


def test_docx_round_trip(tmp_path, sample_docx_file):
    """Validates round-trip: DOCX -> Document Model -> DOCX -> Document Model."""
    service = DOCXService()

    # Step 1: Parse original
    doc_model1 = service.parse_to_document(sample_docx_file)

    # Step 2: Export
    exported_path = tmp_path / "round_trip.docx"
    service.export_to_docx(doc_model1, exported_path)
    assert exported_path.exists()

    # Step 3: Parse exported
    doc_model2 = service.parse_to_document(exported_path)

    full_text1 = doc_model1.get_full_text()
    full_text2 = doc_model2.get_full_text()

    assert "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in full_text2
    assert "Phát triển phần mềm Document Assistant" in full_text2
