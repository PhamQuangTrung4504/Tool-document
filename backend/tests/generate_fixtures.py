"""Generates realistic test fixtures for real-world document auditing.

Creates genuine Vietnamese administrative documents, multi-page PDFs, scanned PDFs,
landscape documents, tables, rotated images, and noisy scans.
"""

from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).parent.parent))

import cv2
import numpy as np
import pymupdf
from PIL import Image, ImageDraw, ImageFont
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

from app.services.pdf.pdf_service import get_system_font_path


def create_all_fixtures(fixtures_dir: Path) -> dict:
    """Generates all comprehensive test fixtures and returns mapping of paths."""
    fixtures_dir.mkdir(parents=True, exist_ok=True)
    font_path = get_system_font_path() or "C:/Windows/Fonts/arial.ttf"
    results = {}

    # -------------------------------------------------------------
    # 1. Real Vietnamese Administrative PDF (Portrait, 2 pages)
    # -------------------------------------------------------------
    admin_pdf_path = fixtures_dir / "van_ban_hanh_chinh.pdf"
    doc = pymupdf.open()
    font_kwargs = {"fontfile": font_path, "fontname": "arial"}

    # Page 1: Header, Decision Title, Table
    p1 = doc.new_page(width=595.28, height=841.89)
    # Header Left
    p1.insert_text((50, 60), "ỦY BAN NHÂN DÂN THÀNH PHỐ", fontsize=11, **font_kwargs)
    p1.insert_text((50, 75), "SỞ KHOA HỌC VÀ CÔNG NGHỆ", fontsize=10, **font_kwargs)
    p1.insert_text((50, 90), "Số: 125/QĐ-SKHCN", fontsize=10, **font_kwargs)

    # Header Right
    p1.insert_text((320, 60), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", fontsize=11, **font_kwargs)
    p1.insert_text((345, 75), "Độc lập - Tự do - Hạnh phúc", fontsize=10, **font_kwargs)
    p1.insert_text((330, 90), "Hà Nội, ngày 15 tháng 10 năm 2026", fontsize=9, **font_kwargs)

    # Title Centered
    p1.insert_text((220, 140), "QUYẾT ĐỊNH", fontsize=14, **font_kwargs)
    p1.insert_text((150, 160), "Về việc phê duyệt đề tài nghiên cứu Document Assistant", fontsize=11, **font_kwargs)

    # Body
    p1.insert_text((50, 200), "GIÁM ĐỐC SỞ KHOA HỌC VÀ CÔNG NGHỆ", fontsize=11, **font_kwargs)
    p1.insert_text((50, 225), "Căn cứ Luật Công nghệ thông tin ngày 29 tháng 6 năm 2006;", fontsize=10, **font_kwargs)
    p1.insert_text((50, 245), "Xét đề nghị của Trưởng phòng Kế hoạch - Tài chính.", fontsize=10, **font_kwargs)
    p1.insert_text((240, 275), "QUYẾT ĐỊNH:", fontsize=12, **font_kwargs)
    p1.insert_text((50, 305), "Điều 1. Phê duyệt danh mục nhiệm vụ khoa học và công nghệ năm 2026:", fontsize=10, **font_kwargs)

    # Draw Table
    # Header row
    p1.draw_rect(pymupdf.Rect(50, 320, 545, 345), color=(0.2, 0.2, 0.2), width=1)
    p1.insert_text((55, 337), "STT", fontsize=10, **font_kwargs)
    p1.insert_text((90, 337), "Tên đề tài / Dự án", fontsize=10, **font_kwargs)
    p1.insert_text((350, 337), "Chủ nhiệm", fontsize=10, **font_kwargs)
    p1.insert_text((450, 337), "Kinh phí (VNĐ)", fontsize=10, **font_kwargs)

    # Row 1
    p1.draw_rect(pymupdf.Rect(50, 345, 545, 370), color=(0.7, 0.7, 0.7), width=0.5)
    p1.insert_text((55, 362), "01", fontsize=9, **font_kwargs)
    p1.insert_text((90, 362), "Hệ thống Document Engine Offline", fontsize=9, **font_kwargs)
    p1.insert_text((350, 362), "Nguyễn Văn An", fontsize=9, **font_kwargs)
    p1.insert_text((450, 362), "500.000.000", fontsize=9, **font_kwargs)

    # Row 2
    p1.draw_rect(pymupdf.Rect(50, 370, 545, 395), color=(0.7, 0.7, 0.7), width=0.5)
    p1.insert_text((55, 387), "02", fontsize=9, **font_kwargs)
    p1.insert_text((90, 387), "Nhận dạng chữ Việt đa font PaddleOCR", fontsize=9, **font_kwargs)
    p1.insert_text((350, 387), "Trần Thị Bích", fontsize=9, **font_kwargs)
    p1.insert_text((450, 387), "350.000.000", fontsize=9, **font_kwargs)

    # Page 2: Conclusion & Signatures
    p2 = doc.new_page(width=595.28, height=841.89)
    p2.insert_text((50, 60), "Điều 2. Quyết định này có hiệu lực thi hành kể từ ngày ký.", fontsize=10, **font_kwargs)
    p2.insert_text((50, 80), "Điều 3. Chánh Văn phòng, Trưởng các phòng ban chịu trách nhiệm thi hành.", fontsize=10, **font_kwargs)

    p2.insert_text((50, 150), "Nơi nhận:", fontsize=9, **font_kwargs)
    p2.insert_text((50, 165), "- Như Điều 3;", fontsize=8, **font_kwargs)
    p2.insert_text((50, 180), "- Lưu: VT, KHTC.", fontsize=8, **font_kwargs)

    p2.insert_text((380, 150), "GIÁM ĐỐC", fontsize=11, **font_kwargs)
    p2.insert_text((350, 220), "(Đã ký và đóng dấu)", fontsize=9, **font_kwargs)
    p2.insert_text((370, 240), "Lê Hoàng Long", fontsize=11, **font_kwargs)

    doc.save(str(admin_pdf_path))
    doc.close()
    results["admin_pdf"] = admin_pdf_path

    # -------------------------------------------------------------
    # 2. Landscape Financial Table PDF
    # -------------------------------------------------------------
    landscape_path = fixtures_dir / "landscape_table.pdf"
    doc_l = pymupdf.open()
    # A4 Landscape: 841.89 x 595.28
    pl = doc_l.new_page(width=841.89, height=595.28)
    pl.insert_text((280, 50), "BẢNG KÊ KHAI TÀI CHÍNH NĂM 2026", fontsize=14, **font_kwargs)
    pl.draw_rect(pymupdf.Rect(50, 80, 790, 120), color=(0, 0, 0), width=1)
    pl.insert_text((60, 105), "KHOẢN MỤC DOANH THU VÀ CHI PHÍ CHI TIẾT THEO CÁC THÁNG", fontsize=10, **font_kwargs)
    doc_l.save(str(landscape_path))
    doc_l.close()
    results["landscape_pdf"] = landscape_path

    # -------------------------------------------------------------
    # 3. Clean Scanned Image (Vietnamese text with accents)
    # -------------------------------------------------------------
    scan_clean_path = fixtures_dir / "scan_vietnamese_clean.png"
    img_clean = Image.new("RGB", (800, 300), color=(255, 255, 255))
    draw_c = ImageDraw.Draw(img_clean)
    font_pil_large = ImageFont.truetype(font_path, 28)
    font_pil_med = ImageFont.truetype(font_path, 20)
    draw_c.text((40, 40), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", fill=(0, 0, 0), font=font_pil_large)
    draw_c.text((120, 90), "Độc lập - Tự do - Hạnh phúc", fill=(0, 0, 0), font=font_pil_med)
    draw_c.text((40, 160), "Hà Nội, ngày 01 tháng 10 năm 2026", fill=(30, 30, 30), font=font_pil_med)
    draw_c.text((40, 210), "Bảng dấu: ă â ê ô ơ ư đ - Á À Ả Ã Ạ Ắ Ằ Ẳ Ẵ Ặ", fill=(0, 0, 0), font=font_pil_med)
    img_clean.save(str(scan_clean_path))
    results["scan_clean"] = scan_clean_path

    # -------------------------------------------------------------
    # 4. Tilted / Skewed Image (5 degrees skew for deskew test)
    # -------------------------------------------------------------
    scan_tilted_path = fixtures_dir / "scan_tilted_5deg.png"
    tilted_img = img_clean.rotate(-5.0, resample=Image.BICUBIC, expand=True, fillcolor=(255, 255, 255))
    tilted_img.save(str(scan_tilted_path))
    results["scan_tilted"] = scan_tilted_path

    # -------------------------------------------------------------
    # 5. Low Contrast & Noisy Scanned Image
    # -------------------------------------------------------------
    scan_noisy_path = fixtures_dir / "scan_noisy_low_contrast.png"
    arr = np.array(img_clean)
    # Faint text (scale down contrast)
    faint = cv2.convertScaleAbs(arr, alpha=0.5, beta=100)
    # Add Gaussian noise
    noise = np.random.normal(0, 15, faint.shape).astype(np.int16)
    noisy_arr = np.clip(faint.astype(np.int16) + noise, 0, 255).astype(np.uint8)
    cv2.imwrite(str(scan_noisy_path), noisy_arr)
    results["scan_noisy"] = scan_noisy_path

    # -------------------------------------------------------------
    # 6. Scanned PDF (single page image only, no text layer)
    # -------------------------------------------------------------
    scanned_pdf_path = fixtures_dir / "pure_scanned_document.pdf"
    s_doc = pymupdf.open()
    s_page = s_doc.new_page(width=800, height=300)
    s_page.insert_image(pymupdf.Rect(0, 0, 800, 300), filename=str(scan_clean_path))
    s_doc.save(str(scanned_pdf_path))
    s_doc.close()
    results["scanned_pdf"] = scanned_pdf_path

    # -------------------------------------------------------------
    # 7. Mixed PDF (Page 1 text, Page 2 pure scanned image)
    # -------------------------------------------------------------
    mixed_pdf_path = fixtures_dir / "mixed_text_and_scan.pdf"
    m_doc = pymupdf.open()
    # Page 1: Native Text
    mp1 = m_doc.new_page(width=595, height=842)
    mp1.insert_text((50, 80), "TRANG 1: NỘI DUNG VĂN BẢN SỐ HÓA", fontsize=14, **font_kwargs)
    mp1.insert_text((50, 120), "Đây là trang PDF có chứa text layer đầy đủ.", fontsize=11, **font_kwargs)
    # Page 2: Pure scan image
    mp2 = m_doc.new_page(width=800, height=300)
    mp2.insert_image(pymupdf.Rect(0, 0, 800, 300), filename=str(scan_clean_path))
    m_doc.save(str(mixed_pdf_path))
    m_doc.close()
    results["mixed_pdf"] = mixed_pdf_path

    # -------------------------------------------------------------
    # 8. Real DOCX with Vietnamese formatting and table
    # -------------------------------------------------------------
    sample_docx_path = fixtures_dir / "real_contract.docx"
    dx = docx.Document()
    p_title = dx.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("HỢP ĐỒNG CUNG CẤP DỊCH VỤ PHẦN MỀM")
    r_t.bold = True
    r_t.font.size = Pt(14)
    r_t.font.name = "Arial"

    p_body = dx.add_paragraph()
    p_body.paragraph_format.space_before = Pt(6)
    p_body.paragraph_format.space_after = Pt(6)
    r_b = p_body.add_run("Hôm nay, ngày 01 tháng 10 năm 2026, chúng tôi gồm có:")
    r_b.font.size = Pt(11)

    tbl = dx.add_table(rows=3, cols=3)
    tbl.style = "Table Grid"
    tbl.cell(0, 0).text = "Mã Dịch Vụ"
    tbl.cell(0, 1).text = "Tên Gói Dịch Vụ"
    tbl.cell(0, 2).text = "Thời Hạn"
    tbl.cell(1, 0).text = "DV-01"
    tbl.cell(1, 1).text = "Bản quyền Document Assistant Standard"
    tbl.cell(1, 2).text = "12 tháng"
    tbl.cell(2, 0).text = "DV-02"
    tbl.cell(2, 1).text = "Hỗ trợ kỹ thuật On-Premise"
    tbl.cell(2, 2).text = "24 tháng"

    dx.save(str(sample_docx_path))
    results["real_docx"] = sample_docx_path

    # -------------------------------------------------------------
    # 9. Path with Vietnamese Unicode and spaces
    # -------------------------------------------------------------
    unicode_dir = fixtures_dir / "Thư mục kiểm thử tiếng Việt 2026"
    unicode_dir.mkdir(parents=True, exist_ok=True)
    unicode_pdf_path = unicode_dir / "Quyết định phê duyệt số 01.pdf"
    u_doc = pymupdf.open()
    up = u_doc.new_page(width=500, height=400)
    up.insert_text((50, 70), "VĂN BẢN TRONG THƯ MỤC TIẾNG VIỆT", fontsize=12, **font_kwargs)
    u_doc.save(str(unicode_pdf_path))
    u_doc.close()
    results["unicode_path_pdf"] = unicode_pdf_path

    return results


if __name__ == "__main__":
    fixtures = create_all_fixtures(Path("tests/fixtures"))
    print("Created fixtures:")
    for k, v in fixtures.items():
        print(f"  [{k}]: {v}")
