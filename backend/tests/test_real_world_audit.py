"""Comprehensive Real-World Audit and Integration Testing Suite.

Executes real-world document tests against actual files, authentic Vietnamese text,
real PaddleOCR inference, multi-page PDFs, tables, scanned images, and Unicode paths.
"""

from pathlib import Path
import time
import docx
import numpy as np
import pymupdf
import pytest
from PIL import Image

from app.core.exceptions import InvalidDocumentError, UnsupportedFormatError
from app.main import main
from app.models.document import Document, DocumentMetadata, Page
from app.models.elements import BoundingBox, TextElement
from app.services.converter.converter_service import ConversionService
from app.services.docx.docx_service import DOCXService
from app.services.image.preprocessor import ImagePreprocessor
from app.services.ocr.ocr_service import OCRService
from app.services.ocr.paddle_provider import PaddleOCRProvider
from app.services.pdf.analyzer import PDFType, analyze_pdf
from app.services.pdf.pdf_service import PDFService
from app.utils.file_type import FileType, detect_file_type
from tests.generate_fixtures import create_all_fixtures


@pytest.fixture(scope="module")
def fixtures(tmp_path_factory) -> dict:
    """Prepares realistic test fixtures once for the entire module."""
    fix_dir = Path("tests/fixtures")
    return create_all_fixtures(fix_dir)


# ==============================================================================
# SECTION 3: REAL OCR TEST
# ==============================================================================

def test_real_paddleocr_vietnamese_recognition(fixtures):
    """Test 3A-D: Tests real PaddleOCR recognition on clean Vietnamese administrative scan."""
    provider = PaddleOCRProvider()
    if not provider.is_available():
        pytest.skip("PaddleOCR not installed in environment")

    service = OCRService(provider=provider)
    elements = service.recognize_image(fixtures["scan_clean"], preprocess=True)

    assert len(elements) > 0
    all_text = " ".join(el.text for el in elements)

    # Check key Vietnamese phrases
    assert any("CỘNG HÒA" in el.text or "HÒA" in el.text for el in elements)
    assert any("VIỆT NAM" in el.text or "NAM" in el.text for el in elements)
    # Check numbers / dates
    assert any("2026" in el.text for el in elements)
    # Check diacritics
    assert any("ă" in el.text or "â" in el.text or "ê" in el.text or "ô" in el.text or "ơ" in el.text or "ư" in el.text or "đ" in el.text or "Bảng" in el.text for el in elements)

    # Check bounding box validity and confidence
    for el in elements:
        assert el.bbox.width > 0
        assert el.bbox.height > 0
        assert 0.0 <= el.confidence <= 1.0


def test_real_image_preprocessing_deskew_and_noise(fixtures):
    """Test 3G-I: Tests deskew on tilted scan and enhancement on noisy/low-contrast scan."""
    # Test Deskew detection
    tilted_img = ImagePreprocessor.load_image(fixtures["scan_tilted"])
    angle = ImagePreprocessor.detect_skew_angle(tilted_img)
    # The image was rotated by 5 degrees
    assert abs(angle) > 1.0  # Angle must be detected

    deskewed = ImagePreprocessor.deskew(tilted_img)
    assert deskewed.shape[0] > 0 and deskewed.shape[1] > 0

    # Test Noisy / Low contrast preprocessing
    noisy_img = ImagePreprocessor.load_image(fixtures["scan_noisy"])
    enhanced = ImagePreprocessor.process(noisy_img, deskew=False, denoise=True, enhance=True)
    assert enhanced is not None
    assert enhanced.shape == noisy_img.shape


# ==============================================================================
# SECTION 4: PDF REAL-WORLD TEST
# ==============================================================================

def test_real_pdf_analysis_and_structure(fixtures):
    """Test 4.1 - 4.9: Validates PDF classification, text layers, landscape, multi-page."""
    # 1. Text-based Administrative PDF
    admin_analysis = analyze_pdf(fixtures["admin_pdf"])
    assert admin_analysis.pdf_type == PDFType.TEXT_BASED
    assert admin_analysis.total_pages == 2

    # 2. Scanned PDF
    scanned_analysis = analyze_pdf(fixtures["scanned_pdf"])
    assert scanned_analysis.pdf_type == PDFType.SCANNED
    assert scanned_analysis.total_pages == 1

    # 3. Mixed PDF
    mixed_analysis = analyze_pdf(fixtures["mixed_pdf"])
    assert mixed_analysis.pdf_type == PDFType.MIXED
    assert mixed_analysis.total_pages == 2
    assert not mixed_analysis.pages[0].is_scanned
    assert mixed_analysis.pages[1].is_scanned

    # 4. Landscape PDF
    pdf_service = PDFService()
    landscape_doc = pdf_service.parse_to_document(fixtures["landscape_pdf"])
    assert landscape_doc.total_pages == 1
    p = landscape_doc.pages[0]
    assert p.is_landscape
    assert p.width > p.height
    assert "BẢNG KÊ KHAI TÀI CHÍNH" in landscape_doc.get_full_text().replace("\xa0", " ")


def test_real_pdf_to_document_model(fixtures):
    """Test 4: Validates parsing real administrative PDF into Document Model."""
    pdf_service = PDFService()
    doc_model = pdf_service.parse_to_document(fixtures["admin_pdf"])

    assert doc_model.total_pages == 2
    full_text = doc_model.get_full_text().replace("\xa0", " ")

    assert "ỦY BAN NHÂN DÂN THÀNH PHỐ" in full_text
    assert "QUYẾT ĐỊNH" in full_text
    assert "GIÁM ĐỐC" in full_text
    assert "Lê Hoàng Long" in full_text

    # Verify elements on Page 1
    p1 = doc_model.pages[0]
    elements = p1.get_text_elements()
    assert len(elements) >= 5

    # Verify alignment detection (e.g. Centered Title)
    centered = [el for el in elements if el.alignment.value == "center"]
    assert len(centered) > 0


# ==============================================================================
# SECTION 5: PDF → DOCX REAL TEST
# ==============================================================================

def test_real_pdf_to_docx_conversion(tmp_path, fixtures):
    """Test 5: Converts administrative PDF to DOCX, verifies file integrity, paragraphs, tables."""
    conv_service = ConversionService()
    out_docx = tmp_path / "van_ban_chuyen_doi.docx"

    t0 = time.perf_counter()
    conv_service.convert(fixtures["admin_pdf"], to_format="docx", output_path=out_docx)
    duration = time.perf_counter() - t0
    print(f"\n[Perf] PDF -> DOCX duration: {duration:.3f}s")

    assert out_docx.exists()
    assert out_docx.stat().st_size > 0

    # Verify DOCX is valid by opening with python-docx
    doc = docx.Document(str(out_docx))
    assert len(doc.paragraphs) > 0

    full_docx_text = " ".join(p.text for p in doc.paragraphs).replace("\xa0", " ")
    assert "ỦY BAN NHÂN DÂN" in full_docx_text
    assert "QUYẾT ĐỊNH" in full_docx_text
    assert "Nguyễn Văn An" in full_docx_text or "Lê Hoàng Long" in full_docx_text


# ==============================================================================
# SECTION 6: IMAGE → DOCX REAL TEST
# ==============================================================================

def test_real_image_to_docx_conversion(tmp_path, fixtures):
    """Test 6: Converts scanned PNG to DOCX via OCR pipeline."""
    conv_service = ConversionService()
    out_docx = tmp_path / "scan_to_docx.docx"

    t0 = time.perf_counter()
    conv_service.convert(fixtures["scan_clean"], to_format="docx", output_path=out_docx)
    duration = time.perf_counter() - t0
    print(f"\n[Perf] Image -> DOCX duration: {duration:.3f}s")

    assert out_docx.exists()
    doc = docx.Document(str(out_docx))
    assert len(doc.paragraphs) > 0
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert len(all_text.strip()) > 0


# ==============================================================================
# SECTION 7: SEARCHABLE PDF REAL TEST
# ==============================================================================

def test_real_searchable_pdf_generation(tmp_path, fixtures):
    """Test 7: Generates searchable PDF from scanned PDF, verifies text layer and searchability."""
    pdf_service = PDFService()
    out_searchable = tmp_path / "searchable_result.pdf"

    pdf_service.create_searchable_pdf(fixtures["scanned_pdf"], out_searchable, lang="vi")
    assert out_searchable.exists()

    with pymupdf.open(str(out_searchable)) as sdoc:
        assert len(sdoc) == 1
        page = sdoc[0]
        extracted_text = page.get_text().replace("\xa0", " ")
        # Verify text layer now exists on previously pure-image PDF
        assert len(extracted_text.strip()) > 0
        assert "CỘNG HÒA" in extracted_text or "VIỆT NAM" in extracted_text or "HÒA" in extracted_text


# ==============================================================================
# SECTION 8: DOCX ROUND TRIP
# ==============================================================================

def test_real_docx_round_trip(tmp_path, fixtures):
    """Test 8: Validates DOCX -> Document Model -> DOCX without content loss."""
    docx_service = DOCXService()

    # Step 1: Parse original real contract
    doc_model = docx_service.parse_to_document(fixtures["real_docx"])
    assert doc_model.total_pages == 1
    assert len(doc_model.pages[0].get_tables()) == 1

    # Step 2: Export to new DOCX
    out_path = tmp_path / "contract_roundtrip.docx"
    docx_service.export_to_docx(doc_model, out_path)
    assert out_path.exists()

    # Step 3: Re-parse exported DOCX
    doc_model2 = docx_service.parse_to_document(out_path)
    full_text1 = doc_model.get_full_text()
    full_text2 = doc_model2.get_full_text()

    assert "HỢP ĐỒNG CUNG CẤP DỊCH VỤ PHẦN MỀM" in full_text2
    assert "Document Assistant Standard" in full_text2
    assert "Hỗ trợ kỹ thuật On-Premise" in full_text2


# ==============================================================================
# SECTION 9: EXPORTER AUDIT (Unicode & Formats)
# ==============================================================================

def test_exporters_unicode_and_formats(tmp_path, fixtures):
    """Test 9: Verifies all exporters (TXT, MD, HTML, PDF, DOCX) produce valid UTF-8 without Mojibake."""
    conv = ConversionService()

    for fmt in ["txt", "md", "html", "pdf", "docx"]:
        out_f = tmp_path / f"admin_export.{fmt}"
        conv.convert(fixtures["admin_pdf"], to_format=fmt, output_path=out_f)
        assert out_f.exists()
        assert out_f.stat().st_size > 0

        # For text-based formats, verify UTF-8 content and Vietnamese letters
        if fmt in ["txt", "md", "html"]:
            content = out_f.read_text(encoding="utf-8").replace("\xa0", " ")
            assert "QUYẾT ĐỊNH" in content
            assert "\ufffd" not in content  # No replacement / corrupted characters


# ==============================================================================
# SECTION 10: FILE SAFETY & CORRUPT/UNICODE PATHS
# ==============================================================================

def test_file_safety_and_unicode_paths(tmp_path, fixtures):
    """Test 10: Validates file handling with Vietnamese Unicode paths, missing files, empty files."""
    conv = ConversionService()

    # 1. Real Vietnamese Unicode path with spaces
    unicode_pdf = fixtures["unicode_path_pdf"]
    assert unicode_pdf.exists()
    out_unicode_txt = tmp_path / "kết_quả_thử_nghiệm.txt"
    res = conv.convert(unicode_pdf, to_format="txt", output_path=out_unicode_txt)
    assert res.exists()
    assert "VĂN BẢN TRONG THƯ MỤC TIẾNG VIỆT" in res.read_text(encoding="utf-8").replace("\xa0", " ")

    # 2. Missing file
    missing_file = tmp_path / "khong_ton_tai.pdf"
    with pytest.raises(InvalidDocumentError):
        conv.convert(missing_file, to_format="txt")

    # 3. Empty file (0 bytes)
    empty_file = tmp_path / "empty_doc.pdf"
    empty_file.touch()
    with pytest.raises(InvalidDocumentError):
        conv.convert(empty_file, to_format="txt")

    # 4. Corrupted file
    corrupt_file = tmp_path / "corrupt.pdf"
    corrupt_file.write_bytes(b"%PDF-1.7 corrupt broken data without end")
    with pytest.raises(InvalidDocumentError):
        conv.convert(corrupt_file, to_format="txt")


# ==============================================================================
# SECTION 11: WINDOWS CLI INTEGRATION TEST
# ==============================================================================

def test_windows_cli_real_execution(tmp_path, fixtures):
    """Test 11: Validates CLI subcommands execution on real files."""
    admin_pdf = str(fixtures["admin_pdf"])

    # Info
    assert main(["info", admin_pdf]) == 0

    # Convert to Markdown
    out_md = str(tmp_path / "cli_out.md")
    assert main(["convert", admin_pdf, "--to", "md", "-o", out_md]) == 0
    assert Path(out_md).exists()

    # Merge
    merged_out = str(tmp_path / "cli_merged_admin.pdf")
    assert main(["merge", admin_pdf, admin_pdf, "-o", merged_out]) == 0
    assert Path(merged_out).exists()

    # Split
    split_out = str(tmp_path / "cli_splits")
    assert main(["split", admin_pdf, "--pages", "1", "-o", split_out]) == 0
    assert len(list(Path(split_out).glob("*.pdf"))) == 1
