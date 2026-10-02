"""Unit tests for Phase 5: Exporters, ConversionService, and CLI commands."""

from pathlib import Path
import docx
import numpy as np
import pymupdf
import pytest
from PIL import Image

from app.exporters import get_exporter
from app.main import main
from app.models.document import Document, DocumentMetadata, Page
from app.models.elements import (
    BoundingBox,
    TableCell,
    TableElement,
    TextAlignment,
    TextElement,
)
from app.services.converter.converter_service import ConversionService
from app.services.docx.docx_service import DOCXService
from app.services.ocr.base import BaseOCRProvider
from app.services.ocr.ocr_service import OCRService
from app.services.pdf.pdf_service import PDFService, get_system_font_path


class TestOCRProvider(BaseOCRProvider):
    @property
    def name(self) -> str:
        return "TestOCRProvider"

    def is_available(self) -> bool:
        return True

    def recognize(self, image: np.ndarray, lang=None, confidence_threshold=0.5):
        return [
            TextElement(
                text="HỢP ĐỒNG MUA BÁN",
                bbox=BoundingBox(x1=50, y1=50, x2=300, y2=80),
                confidence=0.98,
                font_size=14.0,
            ),
            TextElement(
                text="Bên A và Bên B đồng ý các điều khoản sau.",
                bbox=BoundingBox(x1=50, y1=100, x2=450, y2=125),
                confidence=0.96,
                font_size=11.0,
            ),
        ]


@pytest.fixture
def mock_ocr_service() -> OCRService:
    return OCRService(provider=TestOCRProvider())


@pytest.fixture
def sample_pdf(tmp_path) -> Path:
    pdf_p = tmp_path / "input.pdf"
    doc = pymupdf.open()
    font_path = get_system_font_path()
    font_kwargs = {"fontfile": font_path, "fontname": "arial"} if font_path else {}
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 72), "HỢP ĐỒNG KINH TẾ", fontsize=14, **font_kwargs)
    page.insert_text((50, 110), "Điều 1: Mục tiêu hợp đồng", fontsize=11, **font_kwargs)
    doc.save(str(pdf_p))
    doc.close()
    return pdf_p


@pytest.fixture
def sample_image(tmp_path) -> Path:
    img_p = tmp_path / "scan.png"
    arr = np.ones((300, 500, 3), dtype=np.uint8) * 255
    pil = Image.fromarray(arr)
    pil.save(str(img_p))
    return img_p


def test_exporters(tmp_path):
    """Validates individual exporters TXT, MD, HTML, DOCX, PDF."""
    doc = Document(metadata=DocumentMetadata(title="Test Exporters"))
    page = Page(page_number=1, width=595, height=842)
    page.add_element(
        TextElement(
            bbox=BoundingBox(x1=50, y1=50, x2=300, y2=80),
            text="Tiêu đề thử nghiệm",
            font_size=16.0,
            bold=True,
        )
    )
    doc.add_page(page)

    for fmt in ["txt", "md", "html", "docx", "pdf"]:
        exporter = get_exporter(fmt)
        out_file = tmp_path / f"export.{fmt}"
        res = exporter.export(doc, out_file)
        assert res.exists()
        assert res.stat().st_size > 0


def test_conversion_pdf_to_docx(tmp_path, sample_pdf, mock_ocr_service):
    """Validates PDF -> DOCX conversion."""
    conv = ConversionService(ocr_service=mock_ocr_service)
    out_docx = tmp_path / "converted.docx"
    conv.convert(sample_pdf, to_format="docx", output_path=out_docx)

    assert out_docx.exists()
    d = docx.Document(str(out_docx))
    text_content = " ".join(p.text for p in d.paragraphs).replace("\xa0", " ")
    assert "HỢP ĐỒNG KINH TẾ" in text_content


def test_conversion_pdf_to_txt_and_html(tmp_path, sample_pdf, mock_ocr_service):
    """Validates PDF -> TXT and PDF -> HTML conversions."""
    conv = ConversionService(ocr_service=mock_ocr_service)

    # PDF -> TXT
    out_txt = tmp_path / "converted.txt"
    conv.convert(sample_pdf, to_format="txt", output_path=out_txt)
    assert out_txt.exists()
    content_txt = out_txt.read_text(encoding="utf-8").replace("\xa0", " ")
    assert "HỢP ĐỒNG KINH TẾ" in content_txt

    # PDF -> HTML
    out_html = tmp_path / "converted.html"
    conv.convert(sample_pdf, to_format="html", output_path=out_html)
    assert out_html.exists()
    content_html = out_html.read_text(encoding="utf-8").replace("\xa0", " ")
    assert "HỢP ĐỒNG KINH TẾ" in content_html


def test_conversion_image_to_docx_and_pdf(tmp_path, sample_image, mock_ocr_service):
    """Validates Image -> DOCX and Image -> PDF using OCR."""
    conv = ConversionService(ocr_service=mock_ocr_service)

    # Image -> DOCX
    out_docx = tmp_path / "ocr_out.docx"
    conv.convert(sample_image, to_format="docx", output_path=out_docx)
    assert out_docx.exists()
    d = docx.Document(str(out_docx))
    all_text = " ".join(p.text for p in d.paragraphs)
    assert "HỢP ĐỒNG MUA BÁN" in all_text

    # Image -> PDF
    out_pdf = tmp_path / "ocr_out.pdf"
    conv.convert(sample_image, to_format="pdf", output_path=out_pdf)
    assert out_pdf.exists()
    with pymupdf.open(str(out_pdf)) as pdoc:
        assert len(pdoc) == 1


def test_cli_info_and_convert(tmp_path, sample_pdf):
    """Validates CLI info and convert commands."""
    # Test 'info'
    exit_code = main(["info", str(sample_pdf)])
    assert exit_code == 0

    # Test 'convert'
    out_txt = tmp_path / "cli_converted.txt"
    exit_code = main(["convert", str(sample_pdf), "--to", "txt", "-o", str(out_txt)])
    assert exit_code == 0
    assert out_txt.exists()
    assert "HỢP ĐỒNG" in out_txt.read_text(encoding="utf-8").replace("\xa0", " ")


def test_cli_merge_and_split(tmp_path, sample_pdf):
    """Validates CLI merge and split commands."""
    merged_pdf = tmp_path / "cli_merged.pdf"
    exit_code = main(["merge", str(sample_pdf), str(sample_pdf), "-o", str(merged_pdf)])
    assert exit_code == 0
    assert merged_pdf.exists()
    with pymupdf.open(str(merged_pdf)) as mdoc:
        assert len(mdoc) == 2

    split_dir = tmp_path / "cli_split"
    exit_code = main(["split", str(sample_pdf), "--pages", "1", "-o", str(split_dir)])
    assert exit_code == 0
    assert split_dir.exists()
    split_files = list(split_dir.glob("*.pdf"))
    assert len(split_files) == 1
